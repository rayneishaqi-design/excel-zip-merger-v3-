import csv
import io
import os
import shutil
import tempfile
import zipfile
from collections import Counter
from pathlib import PurePosixPath

import streamlit as st
from openpyxl import Workbook, load_workbook
import xlrd

st.set_page_config(page_title="Gabung Excel dan CSV — ZIP atau Langsung", page_icon="📁", layout="wide")
st.title("📁 Gabung Excel & CSV — ZIP atau Upload Langsung")
st.write("Upload banyak file ZIP, Excel (.xlsx/.xls), CSV, atau campuran semuanya. Data digabung berdasarkan nama kolom. Kapasitas tergantung server.")
st.warning("File diproses di server. Jangan unggah data rahasia ke layanan publik tanpa izin. Lakukan uji dengan data contoh dahulu.")

EXTS = (".xlsx", ".xls", ".csv")
EXCEL_ROW_LIMIT = 1_048_576
EXCEL_COL_LIMIT = 16_384


def cleaned_header(values):
    counts, result = Counter(), []
    for i, cell in enumerate(values):
        base = str(cell).strip() if cell is not None else ""
        if not base:
            base = f"Kolom_{i+1}"
        counts[base] += 1
        result.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return result


def normalized_value(v):
    if v is None:
        return ""
    return str(v)


def open_csv_reader(f):
    raw = io.TextIOWrapper(f, encoding="utf-8-sig", errors="replace", newline="")
    sample = raw.read(32768)
    raw.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return raw, csv.reader(raw, dialect)


def sheets_from_member(z, info, mode, wanted):
    name = info.filename
    if name.lower().endswith('.csv'):
        yield 'CSV'
        return
    if name.lower().endswith('.xlsx'):
        with z.open(info) as f:
            # ZIP member is temporarily buffered one at a time (never all files).
            book = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
            try:
                names = list(book.sheetnames)
            finally:
                book.close()
    else:
        with z.open(info) as f:
            book = xlrd.open_workbook(file_contents=f.read(), on_demand=True)
            names = book.sheet_names()
            book.release_resources()
    if mode == 'Sheet pertama':
        names = names[:1]
    elif mode == 'Nama sheet tertentu':
        names = [wanted] if wanted in names else []
    yield from names


def iter_sheet(z, info, sheet):
    if info.filename.lower().endswith('.csv'):
        with z.open(info) as f:
            wrapper, reader = open_csv_reader(f)
            with wrapper:
                yield from reader
    elif info.filename.lower().endswith('.xlsx'):
        with z.open(info) as f:
            book = load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
            try:
                tab = book[sheet]
                # Excel tertentu menyimpan metadata dimensi salah (misalnya A1:A1)
                # meskipun ribuan baris sebenarnya ada. Paksa baca isi XML asli.
                tab.reset_dimensions()
                for row in tab.iter_rows(values_only=True):
                    yield row
            finally:
                book.close()
    else:
        with z.open(info) as f:
            book = xlrd.open_workbook(file_contents=f.read(), on_demand=True)
            try:
                tab = book.sheet_by_name(sheet)
                for i in range(tab.nrows):
                    yield tab.row_values(i)
            finally:
                book.release_resources()


def build(zipfile_path, mode, wanted, add_source, progress):
    folder = tempfile.mkdtemp(prefix='gabung_')
    csv_path = os.path.join(folder, 'hasil_gabungan.csv')
    xlsx_path = os.path.join(folder, 'hasil_gabungan.xlsx')
    audit_path = os.path.join(folder, 'laporan.csv')
    problems_path = os.path.join(folder, 'masalah.csv')
    audit, problems, tasks, union = [], [], [], []
    try:
        with zipfile.ZipFile(zipfile_path) as z:
            infos = [f for f in z.infolist() if not f.is_dir() and f.filename.lower().endswith(EXTS) and not PurePosixPath(f.filename).name.startswith(('.', '~$'))]
            if not infos:
                raise ValueError('Tidak ada file XLSX, XLS, atau CSV di ZIP.')
            # Safety check against malicious archive expansions, not a normal file count/size limit.
            for info in infos:
                if info.flag_bits & 0x1:
                    problems.append((info.filename, '-', 'ZIP terenkripsi tidak didukung'))
                    continue
                if info.file_size > 100 * 1024 * 1024 and info.file_size / max(info.compress_size, 1) > 1000:
                    problems.append((info.filename, '-', 'Rasio kompresi mencurigakan; dilewati demi keamanan'))
                    continue
                try:
                    names = list(sheets_from_member(z, info, mode, wanted))
                    if not names:
                        problems.append((info.filename, wanted, 'Sheet tidak ditemukan'))
                    for sheet in names:
                        try:
                            rows = iter_sheet(z, info, sheet)
                            header = cleaned_header(next(rows))
                            if not header:
                                raise ValueError('Header kosong')
                            tasks.append((info, sheet, header))
                            for col in header:
                                if col not in union:
                                    union.append(col)
                        except StopIteration:
                            problems.append((info.filename, sheet, 'Sheet kosong'))
                        except Exception as exc:
                            problems.append((info.filename, sheet, f'{type(exc).__name__}: {str(exc)[:150]}'))
                except Exception as exc:
                    problems.append((info.filename, '-', f'{type(exc).__name__}: {str(exc)[:150]}'))
            if not tasks:
                raise ValueError('Tidak ada sheet yang bisa diproses.')
            headers = (['_sumber_file', '_sumber_sheet'] if add_source else []) + union
            # Keep Excel-compatible sheet limits; CSV has no 1,048,576-row worksheet limit.
            xlsx_supported = len(headers) <= EXCEL_COL_LIMIT
            book = Workbook(write_only=True) if xlsx_supported else None
            sheet_index = 0
            sheet_rows = 0
            def new_excel_sheet():
                nonlocal sheet_index, sheet_rows
                sheet_index += 1
                ws = book.create_sheet(f'Gabungan_{sheet_index}')
                ws.append(headers)
                sheet_rows = 1
                return ws
            ws = new_excel_sheet() if book else None
            total_rows = 0
            with open(csv_path, 'w', newline='', encoding='utf-8-sig') as out:
                writer = csv.writer(out)
                writer.writerow(headers)
                for index, (info, sheet, local_header) in enumerate(tasks):
                    processed, empty = 0, 0
                    try:
                        rows = iter_sheet(z, info, sheet)
                        next(rows)
                        index_map = [headers.index(c) for c in local_header]
                        for values in rows:
                            data = [normalized_value(v) for v in values]
                            if not any(v.strip() for v in data):
                                empty += 1
                                continue
                            record = [''] * len(headers)
                            if add_source:
                                record[0], record[1] = info.filename, sheet
                            for j, value in enumerate(data):
                                if j < len(index_map):
                                    record[index_map[j]] = value
                                elif value.strip():
                                    raise ValueError('Ada nilai di luar jumlah kolom header')
                            writer.writerow(record)
                            if ws:
                                if sheet_rows >= EXCEL_ROW_LIMIT:
                                    ws = new_excel_sheet()
                                ws.append([v if len(v) <= 32767 else v[:32767] for v in record])
                                sheet_rows += 1
                            processed += 1
                            total_rows += 1
                        audit.append((info.filename, sheet, processed, empty, 'OK'))
                    except Exception as exc:
                        audit.append((info.filename, sheet, processed, empty, 'SEBAGIAN / GAGAL'))
                        problems.append((info.filename, sheet, f'Saat baca baris: {type(exc).__name__}: {str(exc)[:150]}'))
                    progress.progress((index + 1) / len(tasks), text=f'File/sheet {index+1}/{len(tasks)} • {total_rows:,} baris')
            if book:
                book.save(xlsx_path)
            with open(audit_path, 'w', newline='', encoding='utf-8-sig') as f:
                w = csv.writer(f)
                w.writerow(['File', 'Sheet', 'Baris hasil', 'Baris kosong dilewati', 'Status'])
                w.writerows(audit)
            with open(problems_path, 'w', newline='', encoding='utf-8-sig') as f:
                w = csv.writer(f)
                w.writerow(['File', 'Sheet', 'Masalah'])
                w.writerows(problems)
            return {'folder': folder, 'csv': csv_path, 'xlsx': xlsx_path if book else None, 'audit':audit_path,
                    'issues':problems_path,'files':len(infos),'sheets':len(tasks),'rows':total_rows,
                    'problems':len(problems), 'columns':len(headers)}
    except Exception:
        # Retain temp data for debugging until Streamlit process clears it.
        raise

def stage_uploads_as_zip(uploads):
    """Buat arsip sementara dari unggahan langsung dan isi ZIP, tanpa memuat semua file ke RAM."""
    tmp = tempfile.NamedTemporaryFile(suffix='.zip', delete=False)
    tmp.close()
    seen = set()
    staged = 0
    notes = []
    try:
        # ZIP_STORED menghindari kompresi ulang yang lambat; sumber dipindahkan bertahap.
        with zipfile.ZipFile(tmp.name, 'w', compression=zipfile.ZIP_STORED, allowZip64=True) as out:
            for up_index, upload in enumerate(uploads, 1):
                fname = PurePosixPath(upload.name.replace('\\', '/')).name
                if fname.lower().endswith('.zip'):
                    upload.seek(0)
                    try:
                        with zipfile.ZipFile(upload) as incoming:
                            for info in incoming.infolist():
                                inner = PurePosixPath(info.filename.replace('\\', '/'))
                                if info.is_dir() or inner.name.startswith(('.', '~$')) or not inner.name.lower().endswith(EXTS):
                                    continue
                                if info.flag_bits & 1:
                                    notes.append(f'{fname}: anggota terenkripsi dilewati: {info.filename}')
                                    continue
                                # Nama diarsipkan memakai awalan sumber agar asal file terlacak.
                                member = f'Upload_{up_index}_{fname}/{inner.as_posix().lstrip("/")}'
                                if member in seen:
                                    notes.append(f'Nama file berulang di {fname}: {info.filename}')
                                    continue
                                seen.add(member)
                                with incoming.open(info) as src, out.open(member, 'w', force_zip64=True) as dst:
                                    shutil.copyfileobj(src, dst, length=1024*1024)
                                staged += 1
                    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
                        notes.append(f'{fname}: ZIP gagal dibaca ({type(exc).__name__}: {exc})')
                elif fname.lower().endswith(EXTS):
                    member = f'Upload_langsung/{up_index}_{fname}'
                    upload.seek(0)
                    with out.open(member, 'w', force_zip64=True) as dst:
                        shutil.copyfileobj(upload, dst, length=1024*1024)
                    staged += 1
        if not staged:
            raise ValueError('Tidak ada file Excel/CSV valid dalam unggahan.')
        return tmp.name, staged, notes
    except Exception:
        os.remove(tmp.name)
        raise

uploads = st.file_uploader(
    'Pilih satu atau banyak file ZIP / Excel / CSV',
    type=['zip', 'xlsx', 'xls', 'csv'],
    accept_multiple_files=True,
    help='Anda dapat memilih sekaligus beberapa Excel, CSV, ZIP, atau mencampurnya.'
)
st.caption('Untuk memilih banyak file di Windows, tahan Ctrl sambil klik file, atau seret sekaligus ke area upload.')
a,b = st.columns(2)
with a:
    sheet_mode = st.selectbox('Sheet yang digabung', ['Sheet pertama', 'Semua sheet', 'Nama sheet tertentu'])
with b:
    wanted_sheet = st.text_input('Nama sheet', value='Sheet1', disabled=sheet_mode != 'Nama sheet tertentu')
add_source = st.checkbox('Tambahkan kolom asal file dan sheet', value=True)
st.caption('Tidak ada penghapusan duplikat. Baris yang seluruhnya kosong dilewati; setiap baris lainnya dipertahankan. Kolom disatukan menurut nama header.')

if uploads and st.button('Gabungkan file', type='primary'):
    progress = st.empty()
    try:
        source_path, staged_count, staging_notes = stage_uploads_as_zip(uploads)
        try:
            result = build(source_path, sheet_mode, wanted_sheet.strip(), add_source, progress)
            result['staged_count'] = staged_count
            result['staging_notes'] = staging_notes
            st.session_state['result'] = result
        finally:
            os.remove(source_path)
    except (ValueError, zipfile.BadZipFile, OSError, MemoryError, Exception) as exc:
        st.error(f'Proses tidak selesai: {type(exc).__name__}: {exc}')
        st.session_state.pop('result', None)

if 'result' in st.session_state:
    res = st.session_state['result']
    cols = st.columns(4)
    for col, label, value in zip(cols, ['File', 'Sheet', 'Baris gabungan', 'Masalah'], [res['files'],res['sheets'],res['rows'],res['problems']]):
        col.metric(label, f'{value:,}')
    if res.get('staging_notes'):
        st.warning('Beberapa unggahan tidak ikut: ' + ' | '.join(res['staging_notes'][:8]))
    if res['problems'] or res.get('staging_notes'):
        st.error('Ada masalah pembacaan. Hasil mungkin tidak lengkap. Download dan periksa laporan.')
    else:
        st.success('Seluruh sheet terpilih selesai dibaca tanpa error terdeteksi. Tetap cocokkan sampel invoice dan nominal dengan sumber.')
    st.info('Untuk Excel: unduh XLSX agar kolom langsung terpisah. CSV dengan pemisah koma sebaiknya dibuka melalui Data → From Text/CSV di Excel.')
    if res['xlsx']:
        with open(res['xlsx'], 'rb') as f:
            st.download_button('⬇️ Download Excel (.xlsx)', data=f, file_name='hasil_gabungan.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', width='stretch')
    else:
        st.warning('Kolom melebihi batas format Excel (16.384); hanya CSV tersedia.')
    with open(res['csv'], 'rb') as f:
        st.download_button('⬇️ Download CSV lengkap', data=f, file_name='hasil_gabungan.csv', mime='text/csv', width='stretch')
    with open(res['audit'], 'rb') as f:
        st.download_button('⬇️ Download laporan baris', data=f, file_name='laporan_baris.csv', mime='text/csv')
    with open(res['issues'], 'rb') as f:
        st.download_button('⬇️ Download laporan masalah', data=f, file_name='masalah.csv', mime='text/csv')
    st.caption('Catatan: Teks Excel di atas 32.767 karakter per sel terpotong pada XLSX; gunakan CSV sebagai arsip lengkap. Excel hanya menampilkan maksimal 1.048.576 baris per sheet, sehingga dibuat beberapa sheet otomatis.')
